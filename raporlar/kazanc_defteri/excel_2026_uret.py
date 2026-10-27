import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
import os, sys
SEV = os.environ.get("SEVIYE", "TAM")      # TAM (20-100 bin $) | 2000 | 1000
if SEV == "TAM":
    S = pd.read_pickle("/tmp/s2026.pkl")
    K = [("T4", "T4 (şimdiki)"), ("T7K", "T7K 1,15x + koşullu stop"), ("T6", "T6 1,10x + koşullu stop")]
    BAS0, CIKTI = 50000, "Kazanc_2026_Gun_Hafta_Ay.xlsx"
else:
    _ek = "" if SEV == "2000" else f"_{SEV}"
    D = pd.read_pickle(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", f"gunluk_uzun{_ek}.pkl"))
    S = D[D.index >= "2026-01-01"].rename(columns={"t2": "T4", "t1": "T7KKS", "sr": "T7K", "bot": "T1"})
    K = [("T4", "T4 (şimdiki)"), ("T7KKS", "T7K 1,15x + koşullu stop"), ("T7K", "T7K 1,15x eski stop"), ("T1", "Eski bot T1")]
    if "t6" in D.columns: S = S.rename(columns={"t6": "T6"}); K.append(("T6", "T6 1x + koşullu stop (bilgi)"))
    BAS0, CIKTI = int(SEV), f"Kazanc_2026_Gun_Hafta_Ay_{SEV}.xlsx"
NK = len(K)
F = "Arial"; mavi = Font(name=F, color="0000FF"); siyah = Font(name=F); kalin = Font(name=F, bold=True)
yesil = Font(name=F, color="008000")
bas_fill = PatternFill("solid", fgColor="1F3864"); bas_font = Font(name=F, bold=True, color="FFFFFF")
sari = PatternFill("solid", fgColor="FFFF00")
PCT = '0.00%;-0.00%;-'; USD = '$#,##0;($#,##0);-'
GUN = ["Pzt", "Sal", "Çar", "Per", "Cum", "Cmt", "Paz"]
AY = ["Ocak", "Şubat", "Mart", "Nisan", "Mayıs", "Haziran", "Temmuz", "Ağustos", "Eylül"]
wb = Workbook()

def baslik(ws, row, cols):
    for j, c in enumerate(cols, 1):
        x = ws.cell(row=row, column=j, value=c); x.font = bas_font; x.fill = bas_fill
        x.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 32

# --- Varsayımlar
v = wb.active; v.title = "Varsayımlar"
v["A1"] = "2026 kazanç defteri — varsayımlar"; v["A1"].font = Font(name=F, bold=True, size=13)
v["A3"] = "Başlangıç bakiyesi (1 Ocak 2026, $)"; v["A3"].font = siyah
v["B3"] = BAS0; v["B3"].font = mavi; v["B3"].fill = sari; v["B3"].number_format = USD
v["C3"] = ("← değiştirilebilir (kullanıcının 1-5 milyon TL ≈ 20-100 bin $ aralığından örnek)" if SEV == "TAM" else "← değiştirilebilir (bu sayfadaki getiriler " + SEV + " $ seviyesinin evreniyle ölçüldü)"); v["C3"].font = Font(name=F, italic=True)
notlar = [
 "Kaynak: binance-carry-bot backtest motoru (g33.py), KAR_HARITASI madde 601 / 623 / 628 / 629. Günlük getiriler 'Günlük' sayfasında (mavi = motor çıktısı).",
 "Model: gerçekçi (canlı gibi tek hesap girişi + işlem başı 5 bp + düzeltilmiş kademe dolumu), iki dilim hizasının ortalaması. Evren: " + {"TAM": "TAM (≳ 12.500 $ bakiye)", "2000": "D8 (2.000 $; PDF defterleriyle AYNI seri)", "1000": "D9 = D8 + AVAX (1.000 $, madde 349/630)"}[SEV] + ".",
 "T4 = kod varsayılanı, kaldıraçsız. T7K 1,15x" + (" ve T6 1,10x" if SEV == "TAM" else "") + ": günlük getiri × kaldıraç (T4 ile aynı 6 yıllık maks düşüş, madde 596 yöntemi).",
 "Koşullu stop = CARRY_LONG_24S_PIYASA=0,03 (24 sa long stop yalnız piyasa medyanı ≤ −%3 iken). " + ("T6 serisi bot ritminde (8 sa), T7K serisi 4 sa kontrolle (fark ~3 puan/yıl)." if SEV == "TAM" else "Bot ritminde (8 sa'te bir kontrol). 'Eski stop' = borsada koşulsuz 24 sa stop."),
 "Kâr her gün yeniden yatırılır (bileşik). 2026 ölçümü 1 Ocak – 30 Eylül (Ekim verisi henüz yok).",
 "Bunlar geçmiş verinin simülasyonudur; canlı hesabın gerçek sonucu değildir ve gelecek için garanti vermez.",
]
for i, n in enumerate(notlar):
    v.cell(row=5 + i, column=1, value=n).font = Font(name=F, size=9)
v.column_dimensions["A"].width = 38; v.column_dimensions["B"].width = 14

# --- Günlük
g = wb.create_sheet("Günlük")
cols = ["Tarih", "Gün"] + [f"{a} günlük %" for _, a in K] + [f"{a} bakiye $" for _, a in K] + [f"{a} günlük kâr $" for _, a in K]
baslik(g, 1, cols)
n = len(S)
for i, (t, r) in enumerate(S.iterrows()):
    row = i + 2
    g.cell(row=row, column=1, value=t.to_pydatetime()).number_format = "dd.mm.yyyy"
    g.cell(row=row, column=2, value=GUN[t.weekday()])
    for j, (k, _) in enumerate(K):
        c = g.cell(row=row, column=3 + j, value=float(r[k])); c.font = mavi; c.number_format = PCT
        bc = L(3 + NK + j); rc = L(3 + j)
        prev = "Varsayımlar!$B$3" if row == 2 else f"{bc}{row-1}"
        b = g.cell(row=row, column=3 + NK + j, value=f"={prev}*(1+{rc}{row})"); b.number_format = USD
        p = g.cell(row=row, column=3 + 2 * NK + j, value=f"={bc}{row}-{prev}"); p.number_format = USD
    for c in range(1, 3 + 3 * NK): g.cell(row=row, column=c).font = g.cell(row=row, column=c).font.copy(name=F)
son = n + 1
for j in range(1, 3 + 3 * NK): g.column_dimensions[L(j)].width = 15 if j > 2 else 11
g.freeze_panes = "C2"

def donem(ws, ad, gruplar, etiket_ad):
    baslik(ws, 1, [etiket_ad, "Başlangıç", "Bitiş", "Gün"] + [f"{a} getiri %" for _, a in K] + [f"{a} kâr $" for _, a in K] + [f"{a} dönem sonu bakiye $" for _, a in K])
    for i, (et, r0, r1) in enumerate(gruplar):
        row = i + 2
        ws.cell(row=row, column=1, value=et)
        ws.cell(row=row, column=2, value=f"='Günlük'!A{r0}").number_format = "dd.mm.yyyy"
        ws.cell(row=row, column=3, value=f"='Günlük'!A{r1}").number_format = "dd.mm.yyyy"
        ws.cell(row=row, column=4, value=r1 - r0 + 1)
        for j in range(NK):
            bc = L(3 + NK + j)
            ilk = "Varsayımlar!$B$3" if r0 == 2 else f"'Günlük'!{bc}{r0-1}"
            ws.cell(row=row, column=5 + j, value=f"='Günlük'!{bc}{r1}/{ilk}-1").number_format = '0.0%;-0.0%;-'
            ws.cell(row=row, column=5 + NK + j, value=f"='Günlük'!{bc}{r1}-{ilk}").number_format = USD
            ws.cell(row=row, column=5 + 2 * NK + j, value=f"='Günlük'!{bc}{r1}").number_format = USD
        for c in range(1, 5 + 3 * NK):
            x = ws.cell(row=row, column=c); x.font = Font(name=F, color="008000" if c in (2, 3) or c >= 5 + 2 * NK else "000000")
    for j in range(1, 5 + 3 * NK): ws.column_dimensions[L(j)].width = 14
    ws.column_dimensions["A"].width = 16; ws.freeze_panes = "E2"

idx = list(S.index)
# Haftalık (Pazartesi başlangıçlı ISO hafta)
hw = {}
for i, t in enumerate(idx):
    key = (t - pd.Timedelta(days=t.weekday())); hw.setdefault(key, []).append(i + 2)
haf = [(f"{k:%d.%m} haftası", v_[0], v_[-1]) for k, v_ in hw.items()]
donem(wb.create_sheet("Haftalık"), "Haftalık", haf, "Hafta")
aw = {}
for i, t in enumerate(idx): aw.setdefault(t.month, []).append(i + 2)
aylar = [(f"{AY[m-1]} 2026", v_[0], v_[-1]) for m, v_ in aw.items()]
donem(wb.create_sheet("Aylık"), "Aylık", aylar, "Ay")

# --- Özet
o = wb.create_sheet("Özet", 1)
o["A1"] = "2026 (1 Ocak – 30 Eylül) özet — " + {"TAM": "TAM evren (20-100 bin $)", "2000": "2.000 $ (D8)", "1000": "1.000 $ (D9)"}[SEV]; o["A1"].font = Font(name=F, bold=True, size=13)
baslik(o, 3, ["Ölçü"] + [a for _, a in K])
satir = [("Başlangıç bakiyesi $", lambda j: "=Varsayımlar!$B$3", USD),
         ("30 Eylül bakiyesi $", lambda j: f"='Günlük'!{L(3+NK+j)}{son}", USD),
         ("Toplam kâr $", lambda j: f"='Günlük'!{L(3+NK+j)}{son}-Varsayımlar!$B$3", USD),
         ("Toplam getiri %", lambda j: f"='Günlük'!{L(3+NK+j)}{son}/Varsayımlar!$B$3-1", '0.0%'),
         ("Ortalama günlük getiri %", lambda j: f"=AVERAGE('Günlük'!{L(3+j)}2:{L(3+j)}{son})", PCT),
         ("En iyi gün %", lambda j: f"=MAX('Günlük'!{L(3+j)}2:{L(3+j)}{son})", PCT),
         ("En kötü gün %", lambda j: f"=MIN('Günlük'!{L(3+j)}2:{L(3+j)}{son})", PCT),
         ("Kazançlı gün sayısı", lambda j: f"=COUNTIF('Günlük'!{L(3+j)}2:{L(3+j)}{son},\">0\")", '0'),
         ("Zararlı gün sayısı", lambda j: f"=COUNTIF('Günlük'!{L(3+j)}2:{L(3+j)}{son},\"<0\")", '0'),
         ("−%3'ten kötü gün", lambda j: f"=COUNTIF('Günlük'!{L(3+j)}2:{L(3+j)}{son},\"<-0.03\")", '0'),
         ("−%5'ten kötü gün", lambda j: f"=COUNTIF('Günlük'!{L(3+j)}2:{L(3+j)}{son},\"<-0.05\")", '0'),
         ("En kötü hafta %", lambda j: f"=MIN(Haftalık!{L(5+j)}2:{L(5+j)}{len(haf)+1})", '0.0%'),
         ("En kötü ay %", lambda j: f"=MIN(Aylık!{L(5+j)}2:{L(5+j)}{len(aylar)+1})", '0.0%'),
         ("Zararlı hafta sayısı", lambda j: f"=COUNTIF(Haftalık!{L(5+j)}2:{L(5+j)}{len(haf)+1},\"<0\")", '0')]
for i, (ad, f, fmt) in enumerate(satir):
    row = 4 + i; o.cell(row=row, column=1, value=ad).font = siyah
    for j in range(NK):
        c = o.cell(row=row, column=2 + j, value=f(j)); c.number_format = fmt; c.font = yesil
o.column_dimensions["A"].width = 26
for j in range(2, 2 + NK): o.column_dimensions[L(j)].width = 22
wb.save(CIKTI)
print(len(haf), len(aylar), son)
