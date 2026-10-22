import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
D = pd.read_pickle("gunluk.pkl")
wb = Workbook()
F = "Arial"
bas = Font(name=F, bold=True, color="FFFFFF"); basfill = PatternFill("solid", fgColor="1F4E78")
mavi = Font(name=F, color="0000FF"); siyah = Font(name=F); kalin = Font(name=F, bold=True)
sari = PatternFill("solid", fgColor="FFFF00")
PCT = '0.00%;[Red]-0.00%;-'; USD = '#,##0 "$";[Red]-#,##0 "$";-'

# --- Varsayımlar
v = wb.active; v.title = "Varsayımlar"
rows = [("Başlangıç bakiyesi ($)", 2000, "Kullanıcının sermayesi (~2.000 $)."),
        ("T60 kaldıraç (sağ sütun)", 1.25, "Değiştirilebilir. 1,25x'te tam dönem düşüş ~ -%25 (madde 437 tablosu)."),
        ("Başlangıç tarihi", D.index[0].date(), "Backtest ısınma süresinden sonraki ilk gün."),
        ("Bitiş tarihi", D.index[-1].date(), "Veri sonu (Ekim 2026 funding verisi henüz yayınlanmadı).")]
v["A1"] = "Varsayım"; v["B1"] = "Değer"; v["C1"] = "Not"
for c in "ABC": v[f"{c}1"].font = bas; v[f"{c}1"].fill = basfill
for i, (a, b, c) in enumerate(rows, 2):
    v[f"A{i}"] = a; v[f"B{i}"] = b; v[f"C{i}"] = c
    v[f"A{i}"].font = siyah; v[f"C{i}"].font = siyah; v[f"B{i}"].font = mavi
v["B2"].fill = sari; v["B3"].fill = sari; v["B2"].number_format = USD; v["B3"].number_format = '0.00"x"'
v["B4"].number_format = v["B5"].number_format = "yyyy-mm-dd"
notlar = ["KAYNAK: binance-carry-bot backtest motoru (g33.py), gerçekçi ağırlık kayması modeli (madde 419),",
          "2.000 $ seviyesi (D8 evreni), iki dilim hizasının (FAZ_KAY 0/1) günlük getiri ortalaması.",
          "Eski bot = madde 442 öncesi ayarlar (hacim 100M, yaş 90, kâr al 0,40). T60 = madde 442 ile botun yeni varsayılanı.",
          "T60 x kaldıraç sütunu yaklaşıktır: günlük getiri x kaldıraç (teminat maliyeti ve likidasyon modelde yok).",
          "Bunlar BACKTEST sonuçlarıdır, canlı hesabın gerçek getirisi değildir. Gelecek için garanti değildir.",
          "Mavi = girilen veri / varsayım; siyah = formül. Sarı hücreler değiştirilebilir."]
for i, n in enumerate(notlar, 8):
    v[f"A{i}"] = n; v[f"A{i}"].font = Font(name=F, italic=True)
v.column_dimensions["A"].width = 30; v.column_dimensions["B"].width = 14; v.column_dimensions["C"].width = 75

# --- Günlük
g = wb.create_sheet("Günlük")
h = ["Tarih", "Eski bot %", "Eski bot $", "T60 %", "T60 $", "T60 x kaldıraç %", "T60 x kaldıraç $"]
for j, t in enumerate(h, 1):
    c = g.cell(1, j, t); c.font = bas; c.fill = basfill; c.alignment = Alignment(horizontal="center", wrap_text=True)
for i, (d, r) in enumerate(D.iterrows(), 2):
    g.cell(i, 1, d.date()).number_format = "yyyy-mm-dd"
    g.cell(i, 2, round(float(r.bot), 8)).font = mavi
    g.cell(i, 4, round(float(r.t60), 8)).font = mavi
    g.cell(i, 6, f"=D{i}*Varsayımlar!$B$3")
    prev = "Varsayımlar!$B$2" if i == 2 else None
    g.cell(i, 3, f"={prev or f'C{i-1}'}*(1+B{i})")
    g.cell(i, 5, f"={prev or f'E{i-1}'}*(1+D{i})")
    g.cell(i, 7, f"={prev or f'G{i-1}'}*(1+F{i})")
    for j in (2, 4, 6): g.cell(i, j).number_format = PCT
    for j in (3, 5, 7): g.cell(i, j).number_format = USD
    for j in (1, 3, 5, 6, 7): g.cell(i, j).font = siyah if j != 1 else siyah
N = len(D) + 1
g.freeze_panes = "B2"
for j, w in enumerate([12, 13, 15, 11, 15, 15, 17], 1): g.column_dimensions[get_column_letter(j)].width = w

def ozet(ad, donemler, etiket):
    s = wb.create_sheet(ad)
    hh = [etiket, "Dönem sonu", "Gün", "Eski bot %", "Eski bot $", "T60 %", "T60 $", "T60 x kaldıraç %", "T60 x kaldıraç $",
          "Kayıp gün (eski)", "Kayıp gün (T60)", "En kötü gün (T60)"]
    for j, t in enumerate(hh, 1):
        c = s.cell(1, j, t); c.font = bas; c.fill = basfill; c.alignment = Alignment(horizontal="center", wrap_text=True)
    for i, (et, son, bas_d) in enumerate(donemler, 2):
        s.cell(i, 1, et); s.cell(i, 2, son).number_format = "yyyy-mm-dd"
        rng = lambda col: f"Günlük!${col}$2:${col}${N}"
        s.cell(i, 3, f'=COUNTIFS({rng("A")},">="&DATE({bas_d.year},{bas_d.month},{bas_d.day}),{rng("A")},"<="&B{i})')
        for cb, cp, cd in ((4, 5, "C"), (6, 7, "E"), (8, 9, "G")):
            s.cell(i, cp, f"=INDEX({rng(cd)},MATCH(B{i},{rng('A')},1))")
            onceki = "Varsayımlar!$B$2" if i == 2 else f"{get_column_letter(cp)}{i-1}"
            s.cell(i, cb, f"={get_column_letter(cp)}{i}/{onceki}-1")
            s.cell(i, cb).number_format = PCT; s.cell(i, cp).number_format = USD
        tarih_k = f'{rng("A")},">="&DATE({bas_d.year},{bas_d.month},{bas_d.day}),{rng("A")},"<="&B{i}'
        s.cell(i, 10, f'=COUNTIFS({tarih_k},{rng("B")},"<0")')
        s.cell(i, 11, f'=COUNTIFS({tarih_k},{rng("D")},"<0")')
        s.cell(i, 12, f'=_xlfn.MINIFS({rng("D")},{tarih_k})'); s.cell(i, 12).number_format = PCT
        for j in range(1, 13): s.cell(i, j).font = siyah
    s.freeze_panes = "B2"
    for j, w in enumerate([12, 12, 6, 13, 14, 11, 14, 15, 16, 11, 11, 13], 1): s.column_dimensions[get_column_letter(j)].width = w
    return s, len(donemler) + 1

idx = D.index
ay = []
for p in sorted(set(idx.to_period("M"))):
    m = idx[idx.to_period("M") == p]
    ay.append((p.strftime("%Y-%m"), m[-1].date(), m[0].date()))
yil = []
for y in sorted(set(idx.year)):
    m = idx[idx.year == y]
    et = str(y) + (" (20 Tem-Ara)" if y == 2022 else " (Oca-Eyl)" if y == 2026 else "")
    yil.append((et, m[-1].date(), m[0].date()))
ys, yn = ozet("Yıllık", yil, "Yıl")
# yıllık özet satırları
r = yn + 2
ys.cell(r, 1, "Tüm dönem").font = kalin
for cb, cp in ((4, 5), (6, 7), (8, 9)):
    L = get_column_letter(cp)
    ys.cell(r, cp, f"={L}{yn}").number_format = USD
    ys.cell(r, cb, f"={L}{yn}/Varsayımlar!$B$2-1").number_format = PCT
ys.cell(r + 1, 1, "Yıllık ortalama (bileşik)").font = kalin
for cb, cp in ((4, 5), (6, 7), (8, 9)):
    L = get_column_letter(cp)
    ys.cell(r + 1, cb, f"=({L}{yn}/Varsayımlar!$B$2)^(365.25/COUNT(Günlük!$A$2:$A${N}))-1").number_format = PCT
ys.cell(r + 3, 1, "Not: 2022 ve 2026 tam yıl değil. Bileşik bakiyeler büyük hesaplarda gerçekçi değildir (likidite).").font = Font(name=F, italic=True)
ozet("Aylık", ay, "Ay")
wb.move_sheet("Yıllık", offset=-3); wb.move_sheet("Aylık", offset=-2)
wb.save("kazanc_tablosu.xlsx")
print("ok", N)
